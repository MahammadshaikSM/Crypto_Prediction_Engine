"""Backtest and tune the 5-signal model on CoinMarketCap daily history.

Reproduces the scoring in prediction_tracker.py on daily data:
  vol_adj_mom = mom7d / (volatility_7d * 100 + 1)     (min-max normalised 0-100)
  mom24h      = 1-day % change                        (normalised)
  btc_rel     = mom7d - BTC mom7d                     (normalised)
  vol_surge   = 24h volume / market cap               (normalised)
  rsi_s       = rsi_score(RSI 14)                     (already 0-100)
  score       = w1*vol_adj + w2*mom24h + w3*btc_rel + w4*vol_surge + w5*rsi_s

Each day the top 10 scores are "picks". A pick is a hit if it is in the real
top-10 gainers H days later (H = 3 and 7). Random baseline = 10 / coins scored.

Weight search (weights sum to 100%):
  1. every combination on a 5% grid
  2. pairwise transfers of 1%, then 0.1%, until nothing improves (hill climb)
Tuning uses TRAIN dates only; TEST dates are never seen during tuning.

Usage:
  python backtest/backtest_weights.py --xlsx "C:/Users/moham/Downloads/cmc_combined_v1.xlsx"
  (or --csv daily.csv if you already extracted the "Daily prices" sheet)
Writes backtest/results.json.
"""
import argparse, csv, itertools, json, math
from pathlib import Path
import numpy as np

HERE = Path(__file__).resolve().parent
FEATURES = ["vol_adj_mom", "mom24h", "btc_rel", "vol_surge", "rsi_s"]
LIVE_WEIGHTS = [22, 28, 25, 10, 15]          # current prediction_tracker.py weights (%)
TRAIN_END = "2025-09-30"                      # tune on <= this date, test after it
HORIZONS = (3, 7)


# ---------- data ----------
def read_rows(args):
    if args.csv:
        with open(args.csv, encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                yield r["date"], r["symbol"], r["id"], r["price"], r["volume"], r["mcap"]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
        it = wb["Daily prices"].iter_rows(values_only=True); next(it)
        for r in it:
            if r[0] is not None:
                yield str(r[0])[:10], r[2], r[4], r[5], r[6], r[7]


def stable_ids(args):
    if not args.xlsx:
        coins = Path(args.csv).with_name("coins.csv")          # written alongside daily.csv
        if not coins.exists():
            return set()
        with open(coins, encoding="utf-8") as fh:
            return {r["id"] for r in csv.DictReader(fh) if r["stable"] == "Yes"}
    import openpyxl
    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    it = wb["Coins"].iter_rows(values_only=True); next(it)
    return {str(r[3]) for r in it if r[0] is not None and r[4] == "Yes"}


def to_float(x):
    try:
        v = float(x)
        return v if math.isfinite(v) else np.nan
    except (TypeError, ValueError):
        return np.nan


def load_matrices(args):
    stables = stable_ids(args)
    rows = [r for r in read_rows(args) if str(r[2]) not in stables]
    dates = sorted({r[0] for r in rows}); ids = sorted({str(r[2]) for r in rows})
    di = {d: i for i, d in enumerate(dates)}; ci = {c: i for i, c in enumerate(ids)}
    P = np.full((len(dates), len(ids)), np.nan); V = P.copy(); M = P.copy(); sym = {}
    for d, s, cid, p, v, m in rows:
        i, j = di[d], ci[str(cid)]
        P[i, j], V[i, j], M[i, j] = to_float(p), to_float(v), to_float(m); sym[str(cid)] = s
    return dates, ids, sym, P, V, M


# ---------- features ----------
def rsi14(P, t):
    w = P[t - 14:t + 1]                                   # 15 closes -> 14 changes
    diff = np.diff(w, axis=0)
    g = np.where(diff > 0, diff, 0).mean(0); l = np.where(diff < 0, -diff, 0).mean(0)
    with np.errstate(divide="ignore", invalid="ignore"):
        r = 100 - 100 / (1 + g / l)
    return np.where(l == 0, 100.0, r)


def rsi_score(r):
    return np.select([r >= 80, r >= 70, r >= 60, r >= 40, r >= 30], [0, 20, 50, 70, 85], 100).astype(float)


def norm(x):
    mn, mx = np.nanmin(x), np.nanmax(x)
    return np.full_like(x, 50.0) if mx - mn <= 0 else (x - mn) / (mx - mn) * 100


def day_features(P, V, M, t, btc_j):
    """Return (valid coin indexes, feature matrix [n x 5]) for day t, like the live tracker."""
    p, p1, p7 = P[t], P[t - 1], P[t - 7]
    window = P[t - 14:t + 1]
    ok = np.isfinite(p) & np.isfinite(p1) & np.isfinite(p7) & (p > 0) & (p1 > 0) & (p7 > 0)
    ok &= np.isfinite(window).all(0) & (V[t] > 0) & (M[t] > 0)
    idx = np.where(ok)[0]
    if len(idx) < 20 or not np.isfinite(P[t, btc_j]) or not np.isfinite(P[t - 7, btc_j]):
        return None, None
    mom24 = (p[idx] / p1[idx] - 1) * 100
    mom7 = (p[idx] / p7[idx] - 1) * 100
    rets = np.diff(P[t - 7:t + 1][:, idx], axis=0) / P[t - 7:t][:, idx]
    vol = np.maximum(rets.std(0), 1e-4)
    btc7 = (P[t, btc_j] / P[t - 7, btc_j] - 1) * 100
    F = np.column_stack([
        norm(mom7 / (vol * 100 + 1)),
        norm(mom24),
        norm(mom7 - btc7),
        norm(V[t, idx] / M[t, idx]),
        rsi_score(rsi14(P, t)[idx]),
    ])
    return idx, F


def build_days(dates, P, V, M, btc_j):
    """Precompute features and forward outcomes for every scorable day."""
    days = []
    for t in range(15, len(dates)):
        idx, F = day_features(P, V, M, t, btc_j)
        if idx is None:
            continue
        out = {}
        for h in HORIZONS:
            if t + h >= len(dates):
                continue
            fwd = P[t + h, idx] / P[t, idx] - 1
            good = np.isfinite(fwd)
            if good.sum() < 20:
                continue
            ret = np.where(good, fwd, -np.inf)
            top = np.argsort(-ret)[:10]
            is_top = np.zeros(len(idx), bool); is_top[top] = True
            out[h] = dict(ret=np.where(good, fwd, np.nan) * 100, top=is_top, n=int(good.sum()))
        days.append(dict(t=t, date=dates[t], idx=idx, F=F, out=out))
    return days


# ---------- evaluation ----------
def evaluate(days, w, h, which):
    """Average hit rate (%), random baseline, positive rate and mean pick return vs universe."""
    w = np.asarray(w, float) / 100.0
    hits, base, pos, edge, n = 0.0, 0.0, 0.0, 0.0, 0
    for d in days:
        if h not in d["out"] or (d["date"] <= TRAIN_END) != (which == "train"):
            continue
        o = d["out"][h]; s = d["F"] @ w
        pick = np.argsort(-s)[:10]
        r = o["ret"][pick]
        hits += o["top"][pick].sum() * 10; base += 1000 / o["n"]
        pos += np.nansum(r > 0) * 10; edge += np.nanmean(r) - np.nanmean(o["ret"]); n += 1
    if not n:
        return None
    return dict(hit=hits / n, base=base / n, pos=pos / n, edge=edge / n, days=n)


def simplex(step):
    k = int(round(100 / step))
    for c in itertools.combinations(range(k + 4), 4):          # stars and bars, 5 parts
        parts = [c[0]] + [c[i] - c[i - 1] - 1 for i in range(1, 4)] + [k + 3 - c[3]]
        yield [p * step for p in parts]


def tune(days, h):
    # train-only data for speed
    train = [d for d in days if d["date"] <= TRAIN_END and h in d["out"]]
    Fs = [d["F"] for d in train]; tops = [d["out"][h]["top"] for d in train]
    def score(w):
        w = np.asarray(w) / 100.0
        return sum(top[np.argsort(-(F @ w))[:10]].sum() for F, top in zip(Fs, tops)) * 10 / len(train)
    best_w, best = None, -1
    for w in simplex(5):
        s = score(w)
        if s > best:
            best, best_w = s, w
    trace = [dict(stage="5% grid", w=best_w, hit=round(best, 3))]
    for step in (1.0, 0.1):
        improved = True
        while improved:
            improved = False
            for i, j in itertools.permutations(range(5), 2):
                if best_w[i] - step < -1e-9:
                    continue
                w = list(best_w); w[i] = round(w[i] - step, 1); w[j] = round(w[j] + step, 1)
                s = score(w)
                if s > best + 1e-9:
                    best, best_w, improved = s, w, True
        trace.append(dict(stage=f"{step}% steps", w=best_w, hit=round(best, 3)))
    return [round(x, 1) for x in best_w], trace


def example_day(days, ids, sym, w, h, date):
    d = next((d for d in days if d["date"] >= date and h in d["out"]), None)
    if not d:
        return None
    s = d["F"] @ (np.asarray(w) / 100); pick = np.argsort(-s)[:10]; o = d["out"][h]
    actual = np.argsort(-np.nan_to_num(o["ret"], nan=-1e9))[:10]
    sy = lambda k: sym[ids[d["idx"][k]]]
    return dict(date=d["date"], horizon=h, hits=int(o["top"][pick].sum()),
                picks=[dict(symbol=sy(k), ret=round(float(o["ret"][k]), 2), hit=bool(o["top"][k])) for k in pick],
                actual_top10=[dict(symbol=sy(k), ret=round(float(o["ret"][k]), 2)) for k in actual])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx"); ap.add_argument("--csv"); ap.add_argument("--example", default="2026-07-01")
    args = ap.parse_args()
    dates, ids, sym, P, V, M = load_matrices(args)
    btc_j = ids.index("1")
    days = build_days(dates, P, V, M, btc_j)
    print(f"{len(ids)} coins, {dates[0]} -> {dates[-1]}, {len(days)} scorable days")
    res = dict(range=[dates[0], dates[-1]], train_end=TRAIN_END, coins=len(ids), horizons={})
    for h in HORIZONS:
        w, trace = tune(days, h)
        live = {k: evaluate(days, LIVE_WEIGHTS, h, k) for k in ("train", "test")}
        tuned = {k: evaluate(days, w, h, k) for k in ("train", "test")}
        res["horizons"][h] = dict(weights=dict(zip(FEATURES, w)), trace=trace, live=live, tuned=tuned,
                                  example_live=example_day(days, ids, sym, LIVE_WEIGHTS, h, args.example),
                                  example_tuned=example_day(days, ids, sym, w, h, args.example))
        f = lambda r: f"hit {r['hit']:.2f}% vs random {r['base']:.2f}% | up {r['pos']:.1f}% | edge {r['edge']:+.2f}pp | {r['days']}d"
        print(f"\n== {h}-day horizon ==\n weights {dict(zip(FEATURES, w))}")
        for k in ("train", "test"):
            print(f" {k:5}  live  {f(live[k])}\n {k:5}  tuned {f(tuned[k])}")
    (HERE / "results.json").write_text(json.dumps(res, indent=1, default=float), encoding="utf-8")
    print("\nwrote", HERE / "results.json")


if __name__ == "__main__":
    main()
